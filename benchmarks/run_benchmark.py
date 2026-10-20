"""Measure VERBATIM's real accuracy (WER) on your machine.

Two modes:
  1) Your own recordings — a folder with media files and a human-checked transcript for
     each, same name + .txt  (e.g. "Interview 3.mp4" + "Interview 3.txt"):
        python benchmarks/run_benchmark.py --folder "~/Desktop/verified samples"
  2) Public Earnings-21 corporate-call excerpts (downloaded from GitHub on first use):
        python benchmarks/run_benchmark.py --earnings21 --minutes 5

Compare engines/settings with --engine/--variant/--no-vad/--enhance/--vocab.
Writes benchmark_results.xlsx with per-file WER and every individual error.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from verbatim.accuracy import wer  # noqa: E402
from verbatim.audio import MEDIA_EXTENSIONS  # noqa: E402
from verbatim.pipeline import transcribe_file  # noqa: E402
from verbatim.settings import Settings  # noqa: E402

E21_FILES = ["4320211", "4346818", "4366893", "4387332", "4359971", "4383161"]


def earnings21(cache: Path, minutes: float) -> list[tuple[Path, str, bool]]:
    repo = cache / "speech-datasets"
    if not repo.exists():
        subprocess.run(["git", "clone", "--depth", "1", "--filter=blob:none", "--sparse",
                        "https://github.com/revdotcom/speech-datasets.git", str(repo)], check=True)
        subprocess.run(["git", "-C", str(repo), "sparse-checkout", "set", "--no-cone",
                        *[f"earnings21/media/{f}.mp3" for f in E21_FILES],
                        *[f"earnings21/transcripts/nlp_references/{f}.nlp" for f in E21_FILES]], check=True)
    from verbatim.platform_info import ffmpeg_path

    out = []
    for fid in E21_FILES:
        src = repo / "earnings21" / "media" / f"{fid}.mp3"
        clip = cache / f"earnings21_{fid}_{int(minutes)}min.mp3"
        if not clip.exists():
            subprocess.run([ffmpeg_path(), "-nostdin", "-loglevel", "error", "-y", "-t", str(minutes * 60), "-i",
                            str(src), "-c", "copy", str(clip)], check=True)
        rows = (repo / "earnings21" / "transcripts" / "nlp_references" / f"{fid}.nlp").read_text().splitlines()[1:]
        ref = " ".join(r.split("|")[0] + r.split("|")[4] for r in rows)
        out.append((clip, ref, True))  # excerpt: score against the matching prefix of the reference
    return out


def own_folder(folder: Path) -> list[tuple[Path, str, bool]]:
    out = []
    for media in sorted(folder.iterdir()):
        if media.suffix.lower() in MEDIA_EXTENSIONS:
            ref = media.with_suffix(".txt")
            if ref.exists():
                out.append((media, ref.read_text(encoding="utf-8"), False))
            else:
                print(f"skip {media.name}: no {ref.name}")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--folder")
    ap.add_argument("--earnings21", action="store_true")
    ap.add_argument("--minutes", type=float, default=5)
    ap.add_argument("--engine", default="whisper")
    ap.add_argument("--variant", default="large-v3")
    ap.add_argument("--language", default="en")
    ap.add_argument("--no-vad", action="store_true")
    ap.add_argument("--enhance", default="normalize")
    ap.add_argument("--vocab")
    ap.add_argument("--crosscheck", action="store_true")
    ap.add_argument("--fuse", action="store_true")
    ap.add_argument("--keep-fillers", action="store_true")
    ap.add_argument("--out", default="benchmark_results.xlsx")
    a = ap.parse_args()
    cache = Path.home() / ".verbatim-bench"
    cache.mkdir(exist_ok=True)
    tests = own_folder(Path(a.folder).expanduser()) if a.folder else earnings21(cache, a.minutes)
    s = Settings(engine=a.engine, variant=a.variant, language=None if a.language == "auto" else a.language,
                 use_vad=not a.no_vad, enhancement=a.enhance, cross_check=a.crosscheck, fuse=a.fuse,
                 vocabulary=Path(a.vocab).read_text(encoding="utf-8") if a.vocab else "", fuzzy_vocab=bool(a.vocab))
    from openpyxl import Workbook

    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    summ.append(["File", "Reference words", "Substitutions", "Deletions", "Insertions", "WER", "Accuracy",
                 "Two-engine agreement", "Review items", "Seconds"])
    errs = wb.create_sheet("Errors")
    errs.append(["File", "Type", "Reference said", "Transcript has", "Context"])
    tot_e = tot_w = 0
    for media, ref, prefix in tests:
        t0 = time.time()
        r = transcribe_file(media, s)
        w = wer(ref, r.transcript.text, ref_prefix=prefix, drop_fillers=not a.keep_fillers)
        e = w.substitutions + w.deletions + w.insertions
        tot_e += e
        tot_w += w.ref_words
        summ.append([media.name, w.ref_words, w.substitutions, w.deletions, w.insertions, w.wer, w.accuracy,
                     r.quality.get("agreement"), r.quality.get("review_items"), round(time.time() - t0, 1)])
        for x in w.errors:
            errs.append([media.name, x["type"], x["reference"], x["transcript"], x["context"]])
        print(f"{media.name:40s} WER {w.wer:6.2%}  accuracy {w.accuracy:6.2%}")
    overall = tot_e / max(1, tot_w)
    summ.append(["OVERALL", tot_w, "", "", "", overall, 1 - overall])
    for row in summ.iter_rows(min_row=2):
        row[5].number_format = row[6].number_format = row[7].number_format = "0.00%"
    wb.save(a.out)
    print(f"\nOVERALL WER {overall:.2%}  →  word accuracy {1 - overall:.2%}   ({tot_w} words)   report: {a.out}")


if __name__ == "__main__":
    main()
