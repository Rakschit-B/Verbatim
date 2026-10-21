# Copyright © 2026 Infinite Story LLP. All rights reserved.
# VERBATIM is proprietary software of Infinite Story LLP. Unauthorised copying, modification or
# distribution of this file, via any medium, is strictly prohibited.
"""Master Excel Transcription Log.

A cumulative, append-only SQLite ledger (<output>/.verbatim/ledger.sqlite) records every
batch, file, segment, correction, review item and error. After each batch the ledger
is rendered to MASTER_TRANSCRIPTION_LOG.xlsx. Review-queue Status/Notes typed into the
Excel file are read back before it is regenerated, so reviewers never lose their work.
"""
from __future__ import annotations

import json
import platform
import sqlite3
from datetime import datetime
from pathlib import Path

from .exporters import stamp

SCHEMA = """
CREATE TABLE IF NOT EXISTS batches(id INTEGER PRIMARY KEY, started TEXT, finished TEXT, app_version TEXT,
  machine TEXT, settings TEXT, files_total INT, completed INT DEFAULT 0, failed INT DEFAULT 0,
  skipped INT DEFAULT 0, stopped INT DEFAULT 0);
CREATE TABLE IF NOT EXISTS files(id INTEGER PRIMARY KEY, batch_id INT, source_name TEXT, source_path TEXT,
  source_size INT, signature TEXT, settings_fp TEXT, status TEXT, language TEXT, engine TEXT, model TEXT,
  duration REAL, speech_seconds REAL, word_count INT, mean_conf REAL, low_conf_pct REAL, review_items INT,
  agreement REAL, corrections INT, removed INT, speakers INT, rtf REAL, seconds REAL, started TEXT,
  finished TEXT, output_folder TEXT, txt TEXT, docx TEXT, srt TEXT, json TEXT, rms_dbfs REAL, clipping_pct REAL);
CREATE TABLE IF NOT EXISTS segments(file_id INT, idx INT, start REAL, end REAL, speaker TEXT, text TEXT,
  confidence REAL, flags TEXT);
CREATE TABLE IF NOT EXISTS review(id INTEGER PRIMARY KEY, file_id INT, start REAL, end REAL, speaker TEXT,
  reason TEXT, text TEXT, confidence TEXT, context TEXT, status TEXT DEFAULT 'Open', note TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS corrections(file_id INT, time REAL, original TEXT, corrected TEXT, rule TEXT);
CREATE TABLE IF NOT EXISTS errors(batch_id INT, source_name TEXT, source_path TEXT, error TEXT, trace TEXT, time TEXT);
CREATE TABLE IF NOT EXISTS voice_jobs(id INTEGER PRIMARY KEY, time TEXT, source_name TEXT, transcript_json TEXT,
  src_lang TEXT, tgt_lang TEXT, engine TEXT, model TEXT, voices TEXT, lines INT, edited INT, faster INT, too_long INT,
  voice_wav TEXT, video TEXT, srt TEXT);
CREATE TABLE IF NOT EXISTS music(id INTEGER PRIMARY KEY, time TEXT, title TEXT, take INT, seed INT, bpm TEXT,
  musical_key TEXT, seconds REAL, style TEXT, mode TEXT, wav TEXT, mp3 TEXT, with_voice TEXT);
CREATE INDEX IF NOT EXISTS ix_files_sig ON files(signature, settings_fp);
CREATE INDEX IF NOT EXISTS ix_seg_file ON segments(file_id);
"""

HEAD = "27233F"
ACCENT = "913A80"


class MasterLog:
    FILENAME = "MASTER_TRANSCRIPTION_LOG.xlsx"

    def __init__(self, root: str | Path):
        self.root = Path(root)
        d = self.root / ".verbatim"
        d.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(d / "ledger.sqlite")
        self.db.executescript(SCHEMA)
        self.db.commit()

    # ---------------- recording ----------------
    def start_batch(self, settings, total: int, version: str) -> int:
        from dataclasses import asdict

        cfg = {k: v for k, v in asdict(settings).items() if k not in ("vocabulary",)}
        cfg["vocabulary_terms"] = len([x for x in settings.vocabulary.replace(",", "\n").splitlines() if x.strip()])
        cur = self.db.execute(
            "INSERT INTO batches(started, app_version, machine, settings, files_total) VALUES(?,?,?,?,?)",
            (datetime.now().isoformat(timespec="seconds"), version, f"{platform.system()} {platform.machine()} · {platform.node()}",
             json.dumps(cfg), total))
        self.db.commit()
        return cur.lastrowid

    def finish_batch(self, bid: int, completed: int, failed: int, skipped: int, stopped: bool = False):
        self.db.execute("UPDATE batches SET finished=?, completed=?, failed=?, skipped=?, stopped=? WHERE id=?",
                        (datetime.now().isoformat(timespec="seconds"), completed, failed, skipped, int(stopped), bid))
        self.db.commit()

    def completed_output(self, signature: str, fp: str) -> str | None:
        row = self.db.execute("SELECT json FROM files WHERE signature=? AND settings_fp=? AND status='Completed' "
                              "ORDER BY id DESC LIMIT 1", (signature, fp)).fetchone()
        if row and row[0] and Path(row[0]).exists():
            return row[0]
        return None

    def record_file(self, bid: int, res, signature: str, fp: str, settings) -> int:
        t, q, i, o = res.transcript, res.quality, res.info, res.outputs
        cur = self.db.execute(
            """INSERT INTO files(batch_id, source_name, source_path, source_size, signature, settings_fp, status, language,
            engine, model, duration, speech_seconds, word_count, mean_conf, low_conf_pct, review_items, agreement,
            corrections, removed, speakers, rtf, seconds, started, finished, output_folder, txt, docx, srt, json,
            rms_dbfs, clipping_pct) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (bid, res.source.name, str(res.source.resolve()), res.source.stat().st_size, signature, fp, res.status,
             t.language, t.engine, t.model, i.get("duration"), i.get("speech_seconds"), q.get("word_count"),
             q.get("mean_confidence"), q.get("low_confidence_pct"), q.get("review_items"), q.get("agreement"),
             q.get("corrections"), q.get("removed"), i.get("speakers"), i.get("realtime_factor"), res.seconds,
             res.started, res.finished, o.get("folder"), o.get("txt"), o.get("docx"), o.get("srt"), o.get("json"),
             i.get("rms_dbfs"), i.get("clipping_pct")))
        fid = cur.lastrowid
        self.db.executemany("INSERT INTO segments VALUES(?,?,?,?,?,?,?,?)",
                            [(fid, k, s.start, s.end, s.speaker, s.text, s.confidence, "; ".join(s.flags))
                             for k, s in enumerate(t.segments, 1)])
        self.db.executemany("INSERT INTO review(file_id,start,end,speaker,reason,text,confidence,context) VALUES(?,?,?,?,?,?,?,?)",
                            [(fid, r["start"], r["end"], r.get("speaker", ""), r["reason"], r.get("text", ""),
                              str(r.get("confidence", "")), r.get("context", "")) for r in res.review])
        self.db.executemany("INSERT INTO corrections VALUES(?,?,?,?,?)",
                            [(fid, c["time"], c["original"], c["corrected"], c["rule"]) for c in res.corrections])
        self.db.commit()
        return fid

    def record_manual_edit(self, json_path: str, t, edits: list[dict], ranges: list[tuple[float, float]]) -> bool:
        """Transcript corrected in the editor: refresh its segments, log each change, close review items."""
        row = self.db.execute("SELECT id FROM files WHERE json=? ORDER BY id DESC LIMIT 1", (str(json_path),)).fetchone()
        if not row:
            return False
        fid = row[0]
        self.db.execute("DELETE FROM segments WHERE file_id=?", (fid,))
        self.db.executemany("INSERT INTO segments VALUES(?,?,?,?,?,?,?,?)",
                            [(fid, k, s.start, s.end, s.speaker, s.text, s.confidence, "; ".join(s.flags))
                             for k, s in enumerate(t.segments, 1)])
        self.db.executemany("INSERT INTO corrections VALUES(?,?,?,?,?)",
                            [(fid, e["time"], e["original"], e["corrected"], e.get("rule", "manual edit (editor)"))
                             for e in edits])
        for a, b in ranges:
            self.db.execute("UPDATE review SET status='Corrected', note=CASE WHEN note='' THEN 'Corrected in VERBATIM editor' "
                            "ELSE note END WHERE file_id=? AND status='Open' AND start < ? AND end > ?", (fid, b, a))
        self.db.commit()
        return True

    def record_voice_job(self, job) -> None:
        lines = job.lines
        self.db.execute(
            "INSERT INTO voice_jobs(time, source_name, transcript_json, src_lang, tgt_lang, engine, model, voices, lines, "
            "edited, faster, too_long, voice_wav, video, srt) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (datetime.now().isoformat(timespec="seconds"), Path(job.media).name if job.media else Path(job.source_json).stem,
             job.source_json, job.src_lang, job.tgt_lang, job.engine, job.model,
             json.dumps(job.voices, ensure_ascii=False), len(lines), sum(1 for x in lines if x.edited),
             sum(1 for x in lines if x.fit.startswith("faster")), sum(1 for x in lines if x.fit.startswith("too long")),
             job.outputs.get("voice", ""), job.outputs.get("video", ""), job.outputs.get("srt", "")))
        self.db.commit()

    def record_error(self, bid: int, src: Path, error: str, trace: str = ""):
        self.db.execute("INSERT INTO errors VALUES(?,?,?,?,?,?)", (bid, src.name, str(src), error, trace[-4000:],
                                                                   datetime.now().isoformat(timespec="seconds")))
        self.db.commit()

    # ---------------- Excel ----------------
    def _sync_review_edits(self, path: Path):
        """Pull Status / Reviewer note edits made in Excel back into the ledger."""
        if not path.exists():
            return
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            if "Review Queue" not in wb.sheetnames:
                return
            ws = wb["Review Queue"]
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            ix, ist, inote = header.index("Review ID"), header.index("Status"), header.index("Reviewer note")
            upd = [(r[ist] or "Open", r[inote] or "", r[ix]) for r in rows if r and r[ix]]
            self.db.executemany("UPDATE review SET status=?, note=? WHERE id=?", upd)
            self.db.commit()
            wb.close()
        except Exception:
            pass

    def record_music(self, items: list) -> None:
        for it in items:
            f, r = it.get("files", {}), it.get("request", {})
            self.db.execute("INSERT INTO music(time, title, take, seed, bpm, musical_key, seconds, style, mode, wav, mp3, "
                            "with_voice) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                            (it.get("created"), it.get("title"), it.get("take"), it.get("seed"), str(it.get("bpm") or ""),
                             it.get("key") or "", f.get("seconds"), (r.get("style") or r.get("description") or "")[:300],
                             r.get("mode"), f.get("wav"), f.get("mp3", ""), f.get("with_voice", "")))
        self.db.commit()
        self.export_excel()

    def export_excel(self) -> Path:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        target = self.root / self.FILENAME
        self._sync_review_edits(target)
        db = self.db
        wb = Workbook()
        dash = wb.active
        dash.title = "Dashboard"

        def sheet(title, headers, rows, widths=None):
            ws = wb.create_sheet(title)
            ws.append(headers)
            for r in rows:
                ws.append(list(r))
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=HEAD)
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for k, h in enumerate(headers, 1):
                w = (widths or {}).get(h, max(12, min(40, len(h) + 4)))
                ws.column_dimensions[get_column_letter(k)].width = w
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
            return ws

        files = db.execute("""SELECT f.id, f.batch_id, f.source_name, f.source_path, f.status, f.language, f.engine, f.model,
            f.duration, f.speech_seconds, f.word_count, f.mean_conf, f.low_conf_pct, f.review_items,
            (SELECT COUNT(*) FROM review r WHERE r.file_id=f.id AND r.status='Open'), f.agreement, f.corrections,
            f.removed, f.speakers, f.rtf, f.seconds, f.rms_dbfs, f.clipping_pct, f.started, f.finished,
            f.output_folder, f.txt, f.docx, f.signature FROM files f ORDER BY f.id""").fetchall()
        latest = {}
        for r in files:
            latest[r[28]] = r[0]
        frows = []
        for r in files:
            frows.append([r[0], r[1], r[2], r[3], r[4], "Yes" if latest.get(r[28]) == r[0] else "No", r[5],
                          f"{r[6]} · {r[7]}", stamp(r[8] or 0).split(".")[0], round((r[8] or 0) / 60, 2),
                          round((r[9] or 0) / 60, 2), r[10], r[11], (r[12] or 0) / 100 if r[12] is not None else None,
                          r[13], r[14], r[15], r[16], r[17], r[18], r[19], r[20], r[21], r[22], r[23], r[24],
                          r[25], r[26], r[27]])
        fh = ["File ID", "Batch", "Original file name", "Source path", "Status", "Latest run", "Language", "Engine · model",
              "Duration", "Minutes", "Speech minutes", "Words", "Mean confidence", "Low-confidence words %",
              "Review items", "Open review items", "Two-engine agreement", "Vocabulary corrections",
              "Hallucinations removed", "Speakers", "Realtime factor", "Processing seconds", "Loudness dBFS",
              "Clipping %", "Started", "Finished", "Output folder", "TXT", "DOCX"]
        ws_files = sheet("Files", fh, frows, {"Original file name": 38, "Source path": 50, "Output folder": 40,
                                              "TXT": 30, "DOCX": 30, "Engine · model": 34})
        for row in ws_files.iter_rows(min_row=2):
            for col, fmt in ((13, "0.0%"), (14, "0.0%"), (17, "0.0%")):
                row[col - 1].number_format = fmt
            for col in (27, 28, 29):
                c = row[col - 1]
                if c.value:
                    c.hyperlink = Path(c.value).resolve().as_uri()
                    c.font = Font(color="1F5FBF", underline="single")
                    c.value = Path(c.value).name or c.value
        n = max(2, ws_files.max_row)
        ws_files.conditional_formatting.add(f"Q2:Q{n}", CellIsRule(operator="lessThan", formula=["0.95"],
                                            fill=PatternFill("solid", fgColor="FBE3E3")))
        ws_files.conditional_formatting.add(f"P2:P{n}", CellIsRule(operator="greaterThan", formula=["0"],
                                            fill=PatternFill("solid", fgColor="FFF4CC")))

        seg = db.execute("""SELECT s.file_id, f.source_name, s.idx, s.start, s.end, s.speaker, s.text, s.confidence, s.flags
            FROM segments s JOIN files f ON f.id=s.file_id ORDER BY s.file_id, s.idx""").fetchall()
        latest_ids = set(latest.values())
        trows = [[r[0], r[1], r[2], stamp(r[3]), stamp(r[4]), r[5] or "", r[6], r[7], r[8] or ""]
                 for r in seg if r[0] in latest_ids]
        ws_t = sheet("Master Transcript", ["File ID", "Original file name", "Segment", "Start", "End", "Speaker",
                                           "Transcript", "Confidence", "Flags"], trows,
                     {"Original file name": 34, "Transcript": 90, "Flags": 26})
        for row in ws_t.iter_rows(min_row=2):
            row[7].number_format = "0.0%"
            row[6].alignment = Alignment(vertical="top", wrap_text=True)

        rv = db.execute("""SELECT r.id, r.file_id, f.source_name, r.start, r.end, r.speaker, r.reason, r.text, r.confidence,
            r.context, r.status, r.note FROM review r JOIN files f ON f.id=r.file_id ORDER BY r.file_id, r.start""").fetchall()
        rrows = [[r[0], r[1], r[2], stamp(r[3]), stamp(r[4]), r[5], r[6], r[7], r[8], r[9], r[10], r[11]]
                 for r in rv if r[1] in latest_ids]
        ws_r = sheet("Review Queue", ["Review ID", "File ID", "Original file name", "Start", "End", "Speaker", "Reason",
                                      "Text", "Confidence", "Context", "Status", "Reviewer note"], rrows,
                     {"Original file name": 30, "Context": 70, "Text": 30, "Reviewer note": 30})
        dv = DataValidation(type="list", formula1='"Open,Corrected,Confirmed OK,Inaudible"', allow_blank=False)
        ws_r.add_data_validation(dv)
        if ws_r.max_row > 1:
            dv.add(f"K2:K{ws_r.max_row}")

        sheet("Corrections", ["File ID", "Original file name", "Time", "Heard as", "Corrected to", "Rule"],
              [[r[0], r[1], stamp(r[2]), r[3], r[4], r[5]] for r in db.execute(
                  "SELECT c.file_id, f.source_name, c.time, c.original, c.corrected, c.rule FROM corrections c "
                  "JOIN files f ON f.id=c.file_id ORDER BY c.file_id, c.time")], {"Original file name": 34})
        sheet("Batches", ["Batch", "Started", "Finished", "VERBATIM version", "Machine", "Files", "Completed", "Failed",
                          "Skipped (resumed)", "Stopped early", "Settings"],
              [[*r[:9], "Yes" if r[9] else "No", r[10]] for r in db.execute(
                  "SELECT id, started, finished, app_version, machine, files_total, completed, failed, skipped, stopped, "
                  "settings FROM batches ORDER BY id")], {"Settings": 80, "Machine": 30})
        sheet("Errors", ["Batch", "Original file name", "Source path", "Error", "Time", "Technical detail"],
              [[r[0], r[1], r[2], r[3], r[5], r[4]] for r in db.execute("SELECT * FROM errors ORDER BY rowid")],
              {"Error": 60, "Technical detail": 80, "Source path": 50})

        vj = list(db.execute("SELECT time, source_name, src_lang, tgt_lang, engine, model, voices, lines, edited, faster, "
                             "too_long, voice_wav, video, srt FROM voice_jobs ORDER BY id"))
        if vj:
            sheet("Voice Studio", ["Time", "Original file name", "From", "To", "Translation engine", "Model", "Voices",
                                   "Lines", "Edited lines", "Sped up", "Too long", "Voiceover WAV", "Dubbed video",
                                   "Subtitles"], [list(r) for r in vj],
                  {"Original file name": 34, "Voices": 30, "Voiceover WAV": 50, "Dubbed video": 50, "Subtitles": 50})
        mu = list(db.execute("SELECT time, title, take, seed, bpm, musical_key, seconds, style, mode, wav, mp3, with_voice "
                             "FROM music ORDER BY id"))
        if mu:
            sheet("Music Studio", ["Time", "Title", "Take", "Seed", "BPM", "Key", "Seconds", "Style", "Mode", "WAV", "MP3",
                                   "With voiceover"], [list(r) for r in mu],
                  {"Title": 30, "Style": 60, "WAV": 50, "MP3": 50, "With voiceover": 50})
        try:
            from .consent import ConsentRegister

            reg = ConsentRegister()
            if reg.entries:
                sheet("Voice Consent", ["Person", "Status", "Scope", "Signed", "Expires", "Project", "Signed form",
                                        "Form SHA-256", "Recorded", "Recorded by", "Revoked", "Notes"],
                      [[e["person"], reg.status(e), e["scope"], e["signed"], e.get("expires", ""), e.get("project", ""),
                        e["form"], e["form_sha256"], e["recorded"], e.get("recorded_by", ""), e.get("revoked", ""),
                        e.get("notes", "")] for e in reg.entries],
                      {"Person": 26, "Scope": 40, "Signed form": 50, "Form SHA-256": 30, "Notes": 40})
        except Exception:
            pass

        # ---------- Dashboard (live formulas over the Files sheet) ----------
        dash["A1"] = "VERBATIM — Master Transcription Log"
        dash["A1"].font = Font(size=18, bold=True, color=HEAD)
        dash["A2"] = f"Infinite Story LLP · regenerated {datetime.now():%d %b %Y %H:%M} · ledger: .verbatim/ledger.sqlite"
        dash["A2"].font = Font(italic=True, color="777777")
        last = max(2, ws_files.max_row)
        F = lambda col: f"Files!{col}2:{col}{last}"  # noqa: E731
        kpis = [
            ("Files transcribed (latest runs)", f'=COUNTIFS({F("E")},"Completed",{F("F")},"Yes")', "0"),
            ("Hours of media", f'=SUMIFS({F("J")},{F("F")},"Yes")/60', "0.00"),
            ("Words transcribed", f'=SUMIFS({F("L")},{F("F")},"Yes")', "#,##0"),
            ("Average word confidence", f'=IFERROR(AVERAGEIFS({F("M")},{F("F")},"Yes"),"n/a")', "0.0%"),
            ("Average two-engine agreement", f'=IFERROR(AVERAGEIFS({F("Q")},{F("F")},"Yes"),"n/a")', "0.0%"),
            ("Open review items", f'=COUNTIF(\'Review Queue\'!K:K,"Open")', "#,##0"),
            ("Vocabulary corrections", f'=SUMIFS({F("R")},{F("F")},"Yes")', "#,##0"),
            ("Files failed (all batches)", "=COUNTA(Errors!A:A)-1", "0"),
            ("Batches run", "=COUNTA(Batches!A:A)-1", "0"),
        ]
        for k, (label, formula, fmt) in enumerate(kpis, start=4):
            dash.cell(k, 1, label).font = Font(bold=True)
            c = dash.cell(k, 2, formula)
            c.number_format = fmt
            c.font = Font(size=13, bold=True, color=ACCENT)
            c.alignment = Alignment(horizontal="right")
        dash.cell(14, 1, "How to use").font = Font(bold=True, color=HEAD)
        tips = ["Files: one row per transcription run; 'Latest run' = Yes marks the current version of each file.",
                "Review Queue: low-confidence words, engine disagreements, possible missed speech and hallucinations.",
                "Set Status in the Review Queue (Open / Corrected / Confirmed OK / Inaudible); it is kept on regeneration.",
                "Agreement below 95% (red) or open review items (amber) mean the transcript needs a human pass.",
                "Accuracy target: a file is 'review-complete' when its Open review items reach 0."]
        for k, tip in enumerate(tips, start=15):
            dash.cell(k, 1, "• " + tip)
        dash.column_dimensions["A"].width = 44
        dash.column_dimensions["B"].width = 22
        try:
            wb.save(target)
            return target
        except PermissionError:  # file open in Excel (Windows locks it)
            alt = self.root / f"MASTER_TRANSCRIPTION_LOG (updated {datetime.now():%Y-%m-%d %H%M}).xlsx"
            wb.save(alt)
            return alt
