# Copyright © 2026 Infinite Story LLP. All rights reserved.
# VERBATIM is proprietary software of Infinite Story LLP. Unauthorised copying, modification or
# distribution of this file, via any medium, is strictly prohibited.
from pathlib import Path

from openpyxl import load_workbook

from verbatim import pipeline
from verbatim.engines.base import Engine
from verbatim.exporters import stamp, subtitle_cues
from verbatim.model import Segment, Transcript, Word
from verbatim.settings import Settings


class FakeEngine(Engine):
    id, label, model = "fake", "Fake", "fake-1"

    def transcribe(self, audio, chunks, opts, progress=None):
        segs = []
        for s, e in chunks:
            words = [Word(" Hello", s, s + 0.4, 0.95), Word(" exl", s + 0.4, s + 0.8, 0.3), Word(" team.", s + 0.8, s + 1.2, 0.9)]
            segs.append(Segment(s, s + 1.2, "Hello exl team.", words))
        return Transcript(segs, "en", self.id, self.model)


def test_stamp():
    assert stamp(83.456) == "00:01:23.456" and stamp(83.456, ",") == "00:01:23,456"


def test_subtitle_rules():
    words = [Word(f" word{i}", i * 0.3, i * 0.3 + 0.3) for i in range(60)]
    t = Transcript([Segment(0, 18, "x", words, speaker="Speaker 1")])
    cues = subtitle_cues(t)
    assert cues[0][2].startswith("[Speaker 1]") and not cues[1][2].startswith("[")
    for s, e, txt in cues:
        assert all(len(line) <= 42 for line in txt.splitlines()) and len(txt.splitlines()) <= 2 and e - s <= 7.5


def test_batch_exports_master_log_and_resume(tmp_path, monkeypatch):
    src_dir = tmp_path / "in"
    src_dir.mkdir()
    media = Path(__file__).parent / "data" / "jfk.flac"
    a = src_dir / "Client Interview 01.flac"
    a.write_bytes(media.read_bytes())
    sub = src_dir / "day2"
    sub.mkdir()
    b = sub / "Client Interview 01.flac"  # same name, different folder
    b.write_bytes(media.read_bytes() + b"\0")
    monkeypatch.setattr(pipeline, "build", lambda *a, **k: FakeEngine())
    s = Settings(output_dir=str(tmp_path / "out"), cross_check=False, vocabulary="EXL", formats=["txt", "srt", "vtt", "csv", "docx", "md"])
    files = pipeline.collect_media([src_dir])
    assert len(files) == 2
    out = pipeline.BatchRunner(files, s).run()
    assert out["completed"] == 2 and out["failed"] == 0
    folders = sorted(p.name for p in (tmp_path / "out").iterdir() if p.is_dir() and not p.name.startswith("."))
    assert folders == ["Client Interview 01.flac", "Client Interview 01.flac__02"]
    txt = (tmp_path / "out" / "Client Interview 01.flac" / "Client Interview 01.txt").read_text()
    assert "SOURCE FILE: Client Interview 01.flac" in txt and "Hello EXL team." in txt
    wb = load_workbook(out["excel"])
    assert {"Dashboard", "Files", "Master Transcript", "Review Queue", "Corrections", "Batches", "Errors"} <= set(wb.sheetnames)
    assert wb["Files"].max_row == 3 and wb["Corrections"].max_row >= 3
    # reviewer marks an item; it must survive regeneration
    ws = wb["Review Queue"]
    hdr = [c.value for c in ws[1]]
    ws.cell(2, hdr.index("Status") + 1, "Corrected")
    wb.save(out["excel"])
    out2 = pipeline.BatchRunner(files, s).run()          # resume: nothing re-transcribed
    assert out2["skipped"] == 2 and out2["completed"] == 0
    wb2 = load_workbook(out2["excel"])
    statuses = [r[hdr.index("Status")] for r in wb2["Review Queue"].iter_rows(min_row=2, values_only=True)]
    assert "Corrected" in statuses
    assert wb2["Batches"].max_row == 3


def test_failure_is_isolated(tmp_path, monkeypatch):
    bad = tmp_path / "broken.mp4"
    bad.write_bytes(b"not a video")
    good = tmp_path / "good.flac"
    good.write_bytes((Path(__file__).parent / "data" / "jfk.flac").read_bytes())
    monkeypatch.setattr(pipeline, "build", lambda *a, **k: FakeEngine())
    s = Settings(output_dir=str(tmp_path / "out"), cross_check=False)
    out = pipeline.BatchRunner([bad, good], s).run()
    assert out["failed"] == 1 and out["completed"] == 1
    wb = load_workbook(out["excel"])
    assert wb["Errors"].cell(2, 2).value == "broken.mp4"


def test_paragraphs_group_speaker_turns():
    from verbatim.exporters import paragraphs

    segs = [Segment(0, 2, "Hello there.", [Word(" Hello", 0, 1), Word(" there.", 1, 2)], speaker="Speaker 1"),
            Segment(2.5, 4, "How are you?", [Word(" How", 2.5, 3)], speaker="Speaker 1"),
            Segment(4.2, 6, "Fine.", [Word(" Fine.", 4.2, 6)], speaker="Speaker 2")]
    p = paragraphs(Transcript(segs))
    assert [x["speaker"] for x in p] == ["Speaker 1", "Speaker 2"] and p[0]["text"] == "Hello there. How are you?"


def test_filename_hints():
    from verbatim.vocabulary import terms_from_filename

    phrase, terms = terms_from_filename("EXL_Town-Hall_ZAGG_final_v2")
    assert "EXL" in terms and "ZAGG" in terms and "final" not in phrase.lower()


def test_silent_file_with_auto_engine_does_not_crash(tmp_path):
    import numpy as np

    from verbatim.audio import write_wav

    src = write_wav(tmp_path / "silence.wav", np.zeros(16000 * 5, np.float32))
    s = Settings(output_dir=str(tmp_path / "out"), cross_check=False, engine="auto", variant="auto")
    out = pipeline.BatchRunner([src], s).run()
    assert out["completed"] == 1 and out["failed"] == 0
